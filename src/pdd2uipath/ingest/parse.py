"""Parse PDF (pdfplumber) and DOCX (python-docx) PDDs into :class:`Document`.

Anchors: DOCX blocks are referenced by section (``§4.2``); PDF blocks by page and
section (``p.3 §4.2``). Tables keep their rows verbatim.
"""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Any

from pdd2uipath.core.state import sha256_file
from pdd2uipath.ingest.models import Block, BlockKind, Document, Table

_HEADING_RE = re.compile(r"^(?P<num>\d+(?:\.\d+)*)\.?\s+(?P<title>[A-Z][^|]{2,80})$")
_APPENDIX_RE = re.compile(r"^(Appendix\s+[A-Z0-9]+)\s*[–-]\s*(.+)$")
_STEP_HEADING_RE = re.compile(r"^(S\d{2})\s*[–-]\s*(.+)$")
_PAGE_FOOTER_RE = re.compile(r"^Page\s+\d+\s+of\s+\d+$", re.I)


class IngestError(ValueError):
    """The PDD could not be parsed."""


def parse_document(path: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    raise IngestError(f"unsupported PDD format {suffix!r}: expected .docx or .pdf")


# ---------------------------------------------------------------- shared helpers


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\xa0", " ")).strip()


def _heading_of(text: str) -> tuple[str, int] | None:
    """Return ``(ref, level)`` when ``text`` looks like a numbered heading."""
    m = _HEADING_RE.match(text)
    if m:
        num = m.group("num")
        return f"§{num}", num.count(".") + 1
    m = _APPENDIX_RE.match(text)
    if m:
        return f"§{m.group(1).replace(' ', '')}", 1
    return None


class _RefTracker:
    def __init__(self, page_aware: bool) -> None:
        self.page_aware = page_aware
        self.section = "front"
        self.step: str | None = None

    def ref(self, page: int | None) -> str:
        sec = self.section if self.step is None else f"{self.section}/{self.step}"
        if self.page_aware and page is not None:
            return f"p.{page} {sec}"
        return sec

    def see(self, text: str) -> tuple[bool, int]:
        """Update the current section; return ``(is_heading, level)``."""
        h = _heading_of(text)
        if h:
            self.section, level = h
            self.step = None
            return True, level
        m = _STEP_HEADING_RE.match(text)
        if m and len(text) < 90:
            self.step = m.group(1)
            return True, 3
        return False, 0


def _title_from_blocks(blocks: list[Block], fallback: str) -> str:
    for i, b in enumerate(blocks):
        if "process definition document" in b.text.lower():
            for nxt in blocks[i + 1 : i + 4]:
                if nxt.kind != "table" and nxt.text and not nxt.text.upper().startswith("STATUS"):
                    return nxt.text
    return fallback


# ---------------------------------------------------------------- DOCX


def parse_docx(path: Path) -> Document:
    import docx
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise IngestError(f"cannot open DOCX {path}: {exc}") from exc
    tracker = _RefTracker(page_aware=False)
    blocks: list[Block] = []
    tables: list[Table] = []
    index = 0
    last_heading = ""
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, d)
            text = _clean(para.text)
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            level = 0
            kind: BlockKind = "paragraph"
            if style.lower().startswith("heading"):
                kind = "heading"
                m = re.search(r"(\d+)", style)
                level = int(m.group(1)) if m else 1
                tracker.see(text)
                last_heading = text
            else:
                is_h, lvl = tracker.see(text)
                if is_h:
                    kind, level = "heading", lvl
                    last_heading = text
                elif "list" in style.lower():
                    kind = "list"
            blocks.append(
                Block(kind=kind, text=text, ref=tracker.ref(None), index=index, level=level)
            )
            index += 1
        elif tag == "tbl":
            t = DocxTable(child, d)
            rows: list[list[str]] = []
            for r in t.rows:
                cells: list[str] = []
                prev = None
                for c in r.cells:
                    if prev is not None and c._tc is prev:  # merged cells repeat the same element
                        continue
                    prev = c._tc
                    cells.append(_clean(c.text.replace("\n", " / ")))
                rows.append(cells)
            if not rows:
                continue
            tbl = Table(
                ref=tracker.ref(None),
                index=index,
                header=rows[0],
                rows=rows[1:],
                context=last_heading,
            )
            tables.append(tbl)
            blocks.append(
                Block(kind="table", text=f"table {len(tables)}", ref=tbl.ref, index=index)
            )
            index += 1
    core_title: str | None
    try:
        core_title = d.core_properties.title or None
    except Exception:
        core_title = None
    title = _title_from_blocks(blocks, core_title or path.stem)
    return Document(
        path=str(path),
        format="docx",
        title=title,
        sha256=sha256_file(path),
        blocks=blocks,
        tables=tables,
        core_title=core_title,
    )


# ---------------------------------------------------------------- PDF


def _repeated_lines(pages: list[list[str]]) -> set[str]:
    counts: Counter[str] = Counter()
    for lines in pages:
        for line in set(lines):
            counts[line] += 1
    return {line for line, n in counts.items() if n >= 2 and len(pages) >= 2 and len(line) < 120}


def _squash(text: str) -> str:
    return re.sub(r"\s+", "", text)


class _PendingTable:
    def __init__(self, table: Table) -> None:
        self.table = table
        self.cells = {c for row in [table.header, *table.rows] for c in row if c}
        self.rows_squashed = [
            _squash("".join(c.replace(" / ", "") for c in row))
            for row in [table.header, *table.rows]
        ]

    def owns_line(self, line: str) -> bool:
        if line in self.cells:
            return True
        s = _squash(line)
        return any(r and (s == r or (len(s) > 12 and s in r)) for r in self.rows_squashed)


def parse_pdf(path: Path) -> Document:
    import pdfplumber

    try:
        pdf = pdfplumber.open(str(path))
    except Exception as exc:
        raise IngestError(f"cannot open PDF {path}: {exc}") from exc
    with pdf:
        page_lines: list[list[str]] = []
        page_tables: list[list[Any]] = []
        for page in pdf.pages:
            text = page.extract_text() or ""
            page_lines.append([_clean(line) for line in text.splitlines() if _clean(line)])
            try:
                page_tables.append(page.extract_tables() or [])
            except Exception:
                page_tables.append([])
        page_count = len(pdf.pages)
    boiler = _repeated_lines(page_lines)
    tracker = _RefTracker(page_aware=True)
    blocks: list[Block] = []
    tables: list[Table] = []
    index = 0
    last_heading = ""

    def emit(pt: _PendingTable, pno: int) -> None:
        nonlocal index
        tbl = pt.table.model_copy(
            update={"ref": tracker.ref(pno), "index": index, "context": last_heading}
        )
        tables.append(tbl)
        blocks.append(
            Block(kind="table", text=f"table {len(tables)}", ref=tbl.ref, index=index, page=pno)
        )
        index += 1

    for pno, (lines, tbls) in enumerate(zip(page_lines, page_tables, strict=True), start=1):
        pending: list[_PendingTable] = []
        for raw in tbls:
            rows = [[_clean((c or "").replace("\n", " / ")) for c in row] for row in raw if row]
            rows = [r for r in rows if any(r)]
            if rows:
                pending.append(
                    _PendingTable(Table(ref="", index=-1, header=rows[0], rows=rows[1:], page=pno))
                )
        emitted: list[_PendingTable] = []
        for line in lines:
            if line in boiler or _PAGE_FOOTER_RE.match(line):
                continue
            is_h, level = tracker.see(line)
            if is_h:
                last_heading = line
                blocks.append(
                    Block(
                        kind="heading",
                        text=line,
                        ref=tracker.ref(pno),
                        index=index,
                        level=level,
                        page=pno,
                    )
                )
                index += 1
                continue
            hit = next((i for i, pt in enumerate(pending) if pt.owns_line(line)), None)
            if hit is not None:
                for pt in pending[: hit + 1]:
                    emit(pt, pno)
                    emitted.append(pt)
                del pending[: hit + 1]
                continue
            if any(pt.owns_line(line) for pt in emitted):
                continue
            is_list = line.startswith(("•", "-", "–", "")) or bool(re.match(r"^\d+\.\s", line))
            text = line.lstrip("•-– ").strip() if is_list else line
            blocks.append(
                Block(
                    kind="list" if is_list else "paragraph",
                    text=text,
                    ref=tracker.ref(pno),
                    index=index,
                    page=pno,
                )
            )
            index += 1
        for pt in pending:
            emit(pt, pno)
    title = _title_from_blocks(blocks, path.stem)
    return Document(
        path=str(path),
        format="pdf",
        title=title,
        sha256=sha256_file(path),
        blocks=blocks,
        tables=tables,
        page_count=page_count,
    )
